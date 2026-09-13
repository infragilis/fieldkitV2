import io
import threading

import openpyxl
from fastapi.testclient import TestClient

from app.main import app
from app.services import cluster_parser
from app.services.cluster_ansible import generate
from app.core.models import ClusterConfig, ClusterSource, NodeConfig
from app.services.cluster_parser import _matches_alias, split_value
from app.services.cluster_validation import (
    classify_address,
    is_valid_hostname,
    is_valid_ip,
)

client = TestClient(app)


def build_workbook(cells):
    """Build an in-memory .xlsx with sanitized (non-customer) data."""
    workbook = openpyxl.Workbook()
    workbook.active.title = "NodeClusterInfo"
    for sheet, coordinate, value in cells:
        if sheet not in workbook.sheetnames:
            workbook.create_sheet(sheet)
        workbook[sheet][coordinate] = value
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def aff_variant_bytes():
    # Mirrors the single-cluster AFF layout with fake data.
    return build_workbook(
        [
            ("NodeClusterInfo", "A28", "Cluster Name"),
            ("NodeClusterInfo", "B28", "demo-cluster"),
            ("NodeClusterInfo", "A49", "DNS Server"),
            ("NodeClusterInfo", "B49", "10.0.0.10"),
            ("NodeClusterInfo", "C49", "10.0.0.11"),
            ("NodeClusterInfo", "A50", "NTP Server"),
            ("NodeClusterInfo", "B50", "time1.example.com"),
            ("NodeClusterInfo", "C50", "time2.example.com"),
            ("NodeClusterInfo", "A63", "Node Name (Auto filled during setup)"),
            ("NodeClusterInfo", "B63", "demo-cluster-01"),
            ("NodeClusterInfo", "C63", "demo-cluster-02"),
        ]
    )


def mcc_variant_bytes():
    # Mirrors the MetroCluster layout: two cluster-name columns and one
    # comma-separated DNS cell.
    return build_workbook(
        [
            ("NodeClusterInfo", "A39", "Cluster Name"),
            ("NodeClusterInfo", "B39", "site-a-cluster"),
            ("NodeClusterInfo", "D39", "site-b-cluster"),
            ("NodeClusterInfo", "A62", "Node Name (Auto filled during setup)"),
            ("NodeClusterInfo", "B62", "site-a-cluster-01"),
            ("NodeClusterInfo", "C62", "site-a-cluster-02"),
            ("NodeClusterInfo", "D62", "site-b-cluster-01"),
            ("NodeClusterInfo", "E62", "site-b-cluster-02"),
            ("NodeClusterInfo", "A96", "DNS Server"),
            ("NodeClusterInfo", "B96", "10.20.0.1 , 10.20.0.2"),
            ("NodeClusterInfo", "A97", "NTP Server"),
            ("NodeClusterInfo", "B97", "Name: time1.example.com / time2.example.com\nAddress: 10.20.1.1"),
        ]
    )


def parse_bytes(data, filename="workbook.xlsx"):
    return client.post(
        "/api/cluster/parse",
        files={"file": (filename, data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


def test_cluster_import_page_route_and_nav():
    response = client.get("/cluster-import")
    assert response.status_code == 200
    assert "Cluster Import" in response.text
    assert 'id="dropzone"' in response.text

    tools_html = __import__("pathlib").Path("app/static/tools.html").read_text(encoding="utf-8")
    assert 'href="/cluster-import"' in tools_html
    app_js = __import__("pathlib").Path("app/static/app.js").read_text(encoding="utf-8")
    assert "nav_cluster" in app_js
    assert 'href="/cluster-import"' not in app_js


# --- label alias matching ---------------------------------------------------


def test_alias_exact_and_boundary_matching():
    assert _matches_alias("cluster name", "cluster name")
    assert _matches_alias("clustername", "clustername")
    assert _matches_alias("ontap cluster", "ontap cluster")
    assert _matches_alias("node name (auto filled during setup)", "node name")
    assert _matches_alias("dns server", "dns server")


def test_alias_does_not_match_related_labels():
    assert not _matches_alias("dns domain name", "dns server")
    assert not _matches_alias("node serial number", "node name")
    assert not _matches_alias("cluster management ip address", "cluster name")


# --- value splitting / normalization ---------------------------------------


def test_split_value_splits_comma_and_strips_prefixes():
    assert split_value("192.0.2.11 , 192.0.2.12", [",", ";", "\n", " / "]) == ["192.0.2.11", "192.0.2.12"]


def test_split_value_extracts_ntp_name_and_address():
    raw = "Name: ntp1.example.com / ntp2.ntp.example.net\nAddress: 192.0.2.10"
    assert split_value(raw, [",", ";", "\n", " / "]) == [
        "ntp1.example.com",
        "ntp2.ntp.example.net",
        "192.0.2.10",
    ]


# --- address classification -------------------------------------------------


def test_hostname_validation():
    assert is_valid_hostname("core-sw-01")
    assert is_valid_hostname("demo-cluster")
    assert not is_valid_hostname("bad_name")
    assert not is_valid_hostname("bad name")


def test_ip_validation():
    assert is_valid_ip("10.0.0.10")
    assert is_valid_ip("2001:db8::1")
    assert not is_valid_ip("999.1.1.1")
    assert not is_valid_ip("10.3.8.")


def test_classify_address():
    assert classify_address("10.0.0.10") == "ip"
    assert classify_address("time1.example.com") == "hostname"
    assert classify_address("not an address") == "invalid"


# --- parsing / validation ---------------------------------------------------


def test_parse_aff_variant_extracts_config():
    response = parse_bytes(aff_variant_bytes())
    assert response.status_code == 200
    payload = response.json()
    config = payload["config"]
    assert config["cluster_name"] == "demo-cluster"
    assert [node["name"] for node in config["nodes"]] == ["demo-cluster-01", "demo-cluster-02"]
    assert config["dns_servers"] == ["10.0.0.10", "10.0.0.11"]
    assert config["ntp_servers"] == ["time1.example.com", "time2.example.com"]
    assert config["source"]["filename"] == "workbook.xlsx"


def test_parse_mcc_variant_flags_ambiguous_cluster_name():
    response = parse_bytes(mcc_variant_bytes(), filename="mcc.xlsx")
    assert response.status_code == 200
    payload = response.json()
    assert payload["config"]["cluster_name"] == "site-a-cluster"
    assert len(payload["config"]["nodes"]) == 4
    assert payload["config"]["dns_servers"] == ["10.20.0.1", "10.20.0.2"]
    codes = {issue["code"] for issue in payload["issues"]}
    assert "ambiguous" in codes


def test_parse_missing_required_fields():
    response = parse_bytes(build_workbook([("NodeClusterInfo", "A1", "Cluster Name")]))
    assert response.status_code == 200
    codes = {(issue["field"], issue["code"]) for issue in response.json()["issues"]}
    assert ("cluster_name", "missing") in codes
    assert ("nodes", "missing") in codes


def test_parse_rejects_xlsx_uncompressed_size_cap(monkeypatch):
    monkeypatch.setattr(cluster_parser, "MAX_XLSX_UNCOMPRESSED_BYTES", 1)
    response = parse_bytes(aff_variant_bytes())
    assert response.status_code == 400
    assert "uncompressed limit" in response.json()["detail"]


def test_parse_rejects_worksheet_dimension_cap(monkeypatch):
    monkeypatch.setattr(cluster_parser, "MAX_WORKSHEET_ROWS", 1)
    response = parse_bytes(build_workbook([("NodeClusterInfo", "A2", "value")]))
    assert response.status_code == 400
    assert "worksheet" in response.json()["detail"].lower()


def test_generated_ansible_files_are_valid_yaml_and_share_node_variable():
    config = ClusterConfig(cluster_name="demo-cluster", nodes=[NodeConfig(name="node-01")], source=ClusterSource(filename="x.xlsx"))
    result = generate(config)
    yaml = __import__("yaml")
    variables = yaml.safe_load(result.variables.content)
    playbook = yaml.safe_load(result.playbook.content)
    assert variables["nodes"] == ["node-01"]
    assert "cluster_nodes" not in result.playbook.content
    assert isinstance(playbook, list)


def test_parse_timeout_is_bounded_and_releases_worker_slot(monkeypatch):
    entered = threading.Event()
    release = threading.Event()

    def blocked_parse(*args):
        entered.set()
        release.wait(2)

    monkeypatch.setattr(cluster_parser, "MAX_XLSX_UNCOMPRESSED_BYTES", 100 * 1024 * 1024)
    from app.api.routes import cluster_import as route
    monkeypatch.setattr(route, "_PARSE_TIMEOUT_SECONDS", 0.01)
    monkeypatch.setattr(route.service, "parse", blocked_parse)
    response = parse_bytes(aff_variant_bytes())
    assert response.status_code == 408
    assert entered.wait(1)
    assert parse_bytes(aff_variant_bytes()).status_code == 503
    release.set()
    assert route._PARSE_SLOT.acquire(timeout=1)
    route._PARSE_SLOT.release()


def test_parse_deduplicates_and_flags_invalid_addresses():
    data = build_workbook(
        [
            ("NodeClusterInfo", "A28", "Cluster Name"),
            ("NodeClusterInfo", "B28", "demo-cluster"),
            ("NodeClusterInfo", "A63", "Node Name (Auto filled during setup)"),
            ("NodeClusterInfo", "B63", "demo-cluster-01"),
            ("NodeClusterInfo", "A49", "DNS Server"),
            ("NodeClusterInfo", "B49", "10.0.0.10"),
            ("NodeClusterInfo", "C49", "10.0.0.10"),
            ("NodeClusterInfo", "D49", "not a valid ip"),
        ]
    )
    response = parse_bytes(data)
    payload = response.json()
    assert payload["config"]["dns_servers"] == ["10.0.0.10", "not a valid ip"]
    codes = {issue["code"] for issue in payload["issues"]}
    assert "duplicate" in codes
    assert "invalid_address" in codes


# --- error handling ---------------------------------------------------------


def test_parse_rejects_unsupported_extension():
    response = parse_bytes(b"anything", filename="notes.txt")
    assert response.status_code == 400
    assert "Unsupported file type" in response.json()["detail"]


def test_parse_rejects_corrupted_workbook():
    response = parse_bytes(b"this is not a zip file", filename="broken.xlsx")
    assert response.status_code == 400


# --- ansible generation -----------------------------------------------------


def test_generate_vars_is_deterministic_and_references_vault_not_secrets():
    payload = parse_bytes(aff_variant_bytes()).json()
    config = payload["config"]
    first = client.post("/api/cluster/generate", json=config).json()
    second = client.post("/api/cluster/generate", json=config).json()
    assert first["variables"]["content"] == second["variables"]["content"]
    assert first["playbook"]["content"] == second["playbook"]["content"]
    assert first["draft"] is True
    variables = first["variables"]["content"]
    assert "demo-cluster" in variables
    assert "vault_ontap_admin_password" in variables
    assert "CHANGE_ME" in variables
    assert any("management IP" in item for item in first["required_inputs"])


# --- integration: upload -> parse -> review -> generate ---------------------


def test_integration_upload_parse_generate():
    upload = parse_bytes(aff_variant_bytes(), filename="integration.xlsx")
    assert upload.status_code == 200
    parsed = upload.json()
    config = parsed["config"]
    assert config["cluster_name"] == "demo-cluster"

    # Simulate a user correction before generation.
    config["dns_servers"] = ["10.0.0.10"]
    generation = client.post("/api/cluster/generate", json=config)
    assert generation.status_code == 200
    result = generation.json()
    assert result["draft"] is True
    assert "demo-cluster" in result["variables"]["content"]
    assert "{{ cluster_name }}" in result["playbook"]["content"]
    assert "netapp.ontap" in result["playbook"]["content"]
    assert result["required_inputs"]
