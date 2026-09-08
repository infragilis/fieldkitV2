"""Workbook opening and field extraction for cluster workbook import.

This module only *reads* workbook content. It never evaluates formulas beyond
their cached values, never executes macros, and never follows external links.
"""

import io

from openpyxl import load_workbook
from openpyxl.workbook import Workbook

from app.core.models import CellLocation, FieldMatch
from app.services.cluster_mappings import (
    FALLBACK_MAX_COL,
    FALLBACK_MAX_ROW,
    FIELD_MAPPINGS,
    PREFERRED_SHEETS,
)


class WorkbookParseError(Exception):
    """Raised for unsupported, malformed, or password-protected workbooks."""


def normalize_label(text: object) -> str:
    return " ".join(str(text).lower().split())


def _matches_alias(normalized: str, alias: str) -> bool:
    if normalized == alias:
        return True
    if normalized.startswith(alias) and len(normalized) > len(alias):
        return normalized[len(alias)] in " (:["
    return False


def _matching_alias(normalized: str, aliases: list[str]) -> str | None:
    for alias in aliases:
        if _matches_alias(normalized, alias):
            return alias
    return None


def _strip_value_prefix(piece: str) -> str:
    text = piece.strip()
    lowered = text.lower()
    for prefix in ("name:", "address:", "server:", "ntp:", "dns:"):
        if lowered.startswith(prefix):
            text = text[len(prefix) :].lstrip()
            lowered = text.lower()
            break
    return text.strip()


def split_value(raw: object, separators: list[str]) -> list[str]:
    """Split a raw cell value into one or more normalized tokens."""
    if raw is None:
        return []
    text = str(raw)
    for separator in separators:
        if separator:
            text = text.replace(separator, "\n")
    tokens: list[str] = []
    for piece in text.split("\n"):
        token = _strip_value_prefix(piece)
        if token:
            tokens.append(token)
    return tokens


def open_workbook(data: bytes, filename: str) -> Workbook:
    """Open workbook bytes, raising a clear error for unsupported files."""
    try:
        return load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - normalize all openpyxl failures
        raise WorkbookParseError(f"Could not open workbook '{filename}': {exc}") from exc


def detect_template(workbook: Workbook) -> str:
    sheets = set(workbook.sheetnames)
    if "NodeClusterInfo" in sheets:
        return "netapp-configbuilder"
    return "unknown"


def _sheet_names(workbook: Workbook) -> list[str]:
    present = [name for name in PREFERRED_SHEETS if name in workbook.sheetnames]
    if present:
        return present
    return list(workbook.sheetnames)


def _collect_right(cells: list, index: int, span: int, gaps_allowed: bool) -> list[tuple[str, str]]:
    values: list[tuple[str, str]] = []
    for offset in range(1, span + 1):
        position = index + offset
        if position >= len(cells):
            break
        cell = cells[position]
        value = cell.value
        if value is None or str(value).strip() == "":
            if not gaps_allowed:
                break
            continue
        values.append((str(value), cell.coordinate))
    return values


def extract_fields(workbook: Workbook, mappings: dict | None = None) -> list[FieldMatch]:
    """Extract raw field matches. Normalization and validation are separate."""
    mappings = mappings or FIELD_MAPPINGS
    matches: list[FieldMatch] = []
    for sheet_name in _sheet_names(workbook):
        worksheet = workbook[sheet_name]
        max_row = FALLBACK_MAX_ROW if sheet_name not in PREFERRED_SHEETS else worksheet.max_row
        max_col = FALLBACK_MAX_COL if sheet_name not in PREFERRED_SHEETS else worksheet.max_column
        for row in worksheet.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
            cells = list(row)
            for index, cell in enumerate(cells):
                if cell.value is None:
                    continue
                label = normalize_label(cell.value)
                if not label:
                    continue
                for field, config in mappings.items():
                    alias = _matching_alias(label, config["aliases"])
                    if alias is None:
                        continue
                    for raw, coordinate in _collect_right(
                        cells, index, config["span"], config["gaps_allowed"]
                    ):
                        for token in split_value(raw, config["split"]):
                            matches.append(
                                FieldMatch(
                                    field=field,
                                    value=token,
                                    label=str(cell.value).strip(),
                                    alias=alias,
                                    location=CellLocation(sheet=sheet_name, cell=coordinate),
                                )
                            )
    return matches
